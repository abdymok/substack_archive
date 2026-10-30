import csv
from copy import copy
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook

book = Path("1070_posts_with_war_day_fixed.xlsx")
patch = list(csv.DictReader(Path("october_subtitle_patch_2026.csv").open(encoding="utf-8-sig"), delimiter=";"))
wb = load_workbook(book)
aliases = {"date":{"date"}, "day_of_war":{"day_of_war","day","war_day"}, "title":{"title"}, "subtitle":{"subtitle"}, "url":{"url","canonical_url"}}
ws = None
cols = None
for sheet in wb.worksheets:
    for rr in range(1, min(sheet.max_row, 15)+1):
        found = {}
        for cc in range(1, sheet.max_column+1):
            key = str(sheet.cell(rr,cc).value or "").strip().lower().replace(" ","_")
            for field,names in aliases.items():
                if key in names: found[field] = cc
        if len(found)==5:
            ws, cols, header = sheet, found, rr
            break
    if ws: break
if ws is None: raise RuntimeError("Required columns not found")

def normdate(v):
    return v.strftime("%Y-%m-%d") if hasattr(v,"strftime") else str(v or "")[:10]
def normday(v):
    try: return str(int(v))
    except: return str(v or "").strip()
def vals(rr):
    return (normdate(ws.cell(rr,cols["date"]).value), normday(ws.cell(rr,cols["day_of_war"]).value), str(ws.cell(rr,cols["title"]).value or "").strip(), str(ws.cell(rr,cols["url"]).value or "").strip())

filled = 0
added = 0
for item in patch:
    key = (item["date"],item["day_of_war"],item["title"],item["url"])
    hits = [rr for rr in range(header+1,ws.max_row+1) if vals(rr)==key]
    if item["action"] == "add_verified_row":
        if hits:
            if len(hits)!=1 or str(ws.cell(hits[0],cols["subtitle"]).value or "").strip()!=item["subtitle"]: raise RuntimeError("Day 1709 conflict")
            continue
        if any(normday(ws.cell(rr,cols["day_of_war"]).value)=="1709" for rr in range(header+1,ws.max_row+1)): raise RuntimeError("Conflicting Day 1709")
        rr=ws.max_row+1
        for cc in range(1,ws.max_column+1):
            if ws.cell(rr-1,cc).has_style: ws.cell(rr,cc)._style=copy(ws.cell(rr-1,cc)._style)
        for field in ("date","day_of_war","title","subtitle","url"):
            value=item[field]
            if field=="date": value=datetime.strptime(value,"%Y-%m-%d").date()
            if field=="day_of_war": value=int(value)
            ws.cell(rr,cols[field]).value=value
        added+=1
    else:
        if len(hits)!=1: raise RuntimeError(f"Expected one match, got {len(hits)}: {key}")
        cell=ws.cell(hits[0],cols["subtitle"])
        old=str(cell.value or "").strip()
        if item["action"]=="intentional_blank":
            if old: raise RuntimeError("Day 1696 is not blank")
        elif old and old!=item["subtitle"]:
            raise RuntimeError(f"Subtitle conflict: {key}")
        elif not old:
            cell.value=item["subtitle"]
            filled+=1

blanks=[]
rows=0
for rr in range(header+1,ws.max_row+1):
    key=vals(rr)
    if any(key):
        rows+=1
        if not str(ws.cell(rr,cols["subtitle"]).value or "").strip(): blanks.append((rr,key))
wb.save(book)
Path("SUBTITLE_AUDIT_2026-10-30.md").write_text(
    "# Subtitle audit — 2026-10-30\n\n"+
    f"- Worksheet: {ws.title}\n- Data rows: {rows}\n- Subtitles filled: {filled}\n- Day 1709 rows added: {added}\n- Remaining blanks: {len(blanks)}\n\n"+
    "## Remaining blanks\n"+"".join(f"- Row {rr}: {key}\n" for rr,key in blanks),encoding="utf-8")
print(f"rows={rows}; filled={filled}; added={added}; remaining_blanks={len(blanks)}")
